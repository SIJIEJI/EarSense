"""Excel input helpers for PPG data."""

from __future__ import annotations

from pathlib import Path
from collections.abc import Iterator

import numpy as np
import pandas as pd
from openpyxl import load_workbook


CANONICAL_COLUMNS = ("Time_PPG", "PPG_Ch1", "PPG_Ch2", "PPG_Ch3")
ALIASES = {
    "Time_PPG": ("Time_PPG", "PPG_Time", "Time PPG"),
    "PPG_Ch1": ("PPG_Ch1", "PPG Ch1", "PPG1"),
    "PPG_Ch2": ("PPG_Ch2", "PPG Ch2", "PPG2"),
    "PPG_Ch3": ("PPG_Ch3", "PPG Ch3", "PPG3"),
}


def _read_sheet(path: Path, sheet: str) -> np.ndarray:
    frame = pd.read_excel(path, sheet_name=sheet, engine="openpyxl")
    frame.columns = [str(c).replace("\ufeff", "").replace("\u200b", "").strip()
                     for c in frame.columns]
    resolved = []
    for canonical in CANONICAL_COLUMNS:
        match = next((name for name in ALIASES[canonical] if name in frame.columns), None)
        if match is None:
            raise ValueError(
                f"Missing '{canonical}' on sheet '{sheet}'. Available: {list(frame.columns)}"
            )
        resolved.append(match)
    data = frame[resolved].apply(pd.to_numeric, errors="coerce").dropna()
    return data.to_numpy(dtype=float)


def iter_ppg_windows(path: Path, window_seconds: float, sampling_rate: float,
                     overlap: float = 0.0) -> Iterator[np.ndarray]:
    """Yield contiguous [time, green, IR, red] windows across workbook sheets."""
    window_samples = int(round(window_seconds * sampling_rate))
    step = window_samples - int(round(window_samples * overlap))
    if window_samples <= 0 or step <= 0:
        raise ValueError("Invalid window size or overlap")

    workbook = load_workbook(path, read_only=True, data_only=True)
    buffer = np.empty((0, 4), dtype=float)
    for sheet in workbook.sheetnames:
        block = _read_sheet(path, sheet)
        if block.size:
            buffer = np.vstack((buffer, block))
        while len(buffer) >= window_samples:
            yield buffer[:window_samples].copy()
            buffer = buffer[step:]
