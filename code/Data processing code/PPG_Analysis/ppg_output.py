"""Excel output helpers for PPG analysis."""

from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import Workbook


def _excel_value(value):
    return float(value) if value is not None and np.isfinite(value) else None


def save_analysis_results(path: Path, times, metrics, hrv) -> None:
    workbook = Workbook()
    metrics_sheet = workbook.active
    metrics_sheet.title = "PPG metrics"
    metrics_sheet.append([
        "TimePoint_s", "HeartRate_bpm", "BreathingRate_bpm", "SpO2_percent",
        "SpO2_ratio", "BreathAmp_median", "BreathAmp_mean", "BreathAmp_SD",
        "BreathCount",
    ])
    for time, row in zip(times, metrics):
        metrics_sheet.append([_excel_value(time), *[_excel_value(v) for v in row]])

    hrv_sheet = workbook.create_sheet("PPG-derived HRV")
    hrv_sheet.append(["TimePoint_s", "RMSSD_ms", "SDNN_ms"])
    for time, row in zip(times, hrv):
        hrv_sheet.append([_excel_value(time), *[_excel_value(v) for v in row]])
    workbook.save(path)


def save_filtered_signals(path: Path, windows, sampling_rate: float) -> None:
    workbook = Workbook(write_only=True)
    sheet_index = 1
    sheet = workbook.create_sheet(f"Filtered{sheet_index}")
    header = ["Time_s", "Green_heart", "Green_breath", "IR_filtered", "Red_filtered"]
    sheet.append(header)
    row_count = 1
    sample_index = 0

    for _, green_hr, green_br, ir_filtered, red_filtered in windows:
        for values in zip(green_hr, green_br, ir_filtered, red_filtered):
            if row_count >= 1_048_576:
                sheet_index += 1
                sheet = workbook.create_sheet(f"Filtered{sheet_index}")
                sheet.append(header)
                row_count = 1
            sheet.append([sample_index / sampling_rate, *[_excel_value(v) for v in values]])
            sample_index += 1
            row_count += 1
    workbook.save(path)
